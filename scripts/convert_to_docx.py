"""
Utility to convert the IEEE-style markdown paper into a DOCX file.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path("paper/face_object_recognition_paper.md")
TARGET = Path("paper/face_object_recognition_paper.docx")


def add_heading_paragraph(document: Document, text: str, italic: bool = False) -> None:
    """
    Add a heading-styled paragraph with 10 pt bold text.
    """
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    if italic:
        run.italic = True


def add_formatted_paragraph(document: Document, raw: str) -> None:
    """
    Parse simple bold (**text**) and italic (_text_) markers within a line.
    """
    paragraph = document.add_paragraph()
    pos = 0
    while pos < len(raw):
        if raw.startswith("**", pos):
            end = raw.find("**", pos + 2)
            if end == -1:
                end = len(raw)
            run = paragraph.add_run(raw[pos + 2 : end])
            run.bold = True
            pos = end + 2
        elif raw.startswith("_", pos):
            end = raw.find("_", pos + 1)
            if end == -1:
                end = len(raw)
            run = paragraph.add_run(raw[pos + 1 : end])
            run.italic = True
            pos = end + 1
        else:
            run = paragraph.add_run(raw[pos])
            pos += 1


def add_table(document: Document, rows: list[str]) -> None:
    """
    Convert markdown-style table rows into a DOCX table.
    """
    header = [cell.strip() for cell in rows[0].strip("|").split("|")]
    body_rows = [r for r in rows[2:]]
    table = document.add_table(rows=1 + len(body_rows), cols=len(header))
    table.style = "Table Grid"
    for idx, text in enumerate(header):
        table.cell(0, idx).text = text
    for row_idx, row in enumerate(body_rows, start=1):
        cells = [cell.strip() for cell in row.strip("|").split("|")]
        for col_idx, cell_text in enumerate(cells):
            table.cell(row_idx, col_idx).text = cell_text


def convert_markdown_to_docx() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source markdown not found: {SOURCE}")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()

    # Configure Normal style to IEEE defaults (Times New Roman, 10 pt).
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(10)

    for section in document.sections:
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        cols = section._sectPr.xpath("./w:cols")
        if cols:
            cols[0].set(qn("w:num"), "2")
            cols[0].set(qn("w:space"), "720")

    i = 0
    in_code_block = False
    code_buffer: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("%"):
            i += 1
            continue

        if stripped.startswith("# "):
            title = stripped[2:].strip()
            title_paragraph = document.add_paragraph()
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title_paragraph.add_run(title)
            run.font.size = Pt(26)
            run.bold = True
            run.font.name = "Times New Roman"
            i += 1
            continue

        if stripped.startswith("## "):
            add_heading_paragraph(document, stripped[3:].strip())
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading_paragraph(document, stripped[4:].strip(), italic=True)
            i += 1
            continue

        if stripped.startswith("- "):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(document, table_lines)
            continue

        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_buffer = []
            else:
                in_code_block = False
                if code_buffer:
                    for item in code_buffer:
                        placeholder = document.add_paragraph(f"[Figure: {item}]")
                        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(stripped)
            i += 1
            continue

        if stripped.startswith("\\["):
            equation_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("\\]"):
                equation_lines.append(lines[i].strip())
                i += 1
            i += 1  # Skip closing \]
            equation_paragraph = document.add_paragraph(" ".join(equation_lines))
            equation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in equation_paragraph.runs:
                run.italic = True
            continue

        if stripped.startswith("[") and stripped[1:2].isdigit():
            document.add_paragraph(stripped)
            i += 1
            continue

        if stripped.startswith("**") or "_" in stripped:
            add_formatted_paragraph(document, line)
            i += 1
            continue

        document.add_paragraph(line.strip())
        i += 1

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)


if __name__ == "__main__":
    convert_markdown_to_docx()


